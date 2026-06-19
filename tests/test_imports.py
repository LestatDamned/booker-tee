from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from fastapi import UploadFile
from openpyxl import Workbook

from app.features.imports.application.documents.management import document_has_linked_operations
from app.features.imports.application.documents.upload import validate_statement_upload
from app.features.imports.application.review.status import raw_transaction_status_for_review_action
from app.features.imports.domain.deduplication import (
    mark_raw_transaction_duplicate,
    possible_duplicate_fingerprint,
)
from app.features.imports.errors import UploadValidationError
from app.features.imports.infrastructure.extraction.openpyxl_extractor import (
    OpenPyxlStatementExtractor,
)
from app.features.imports.infrastructure.extraction.pdfplumber_extractor import (
    PdfPlumberExtractor,
)
from app.features.imports.infrastructure.extraction.resolver import StatementExtractorResolver
from app.features.imports.infrastructure.storage import (
    UploadStorage,
    sanitize_filename,
    sanitize_upload_filename,
)
from app.features.imports.models import RawTransaction, RawTransactionStatus
from app.features.imports.parsing.support.normalization import (
    parse_bank_date,
)
from app.features.imports.presentation.review import (
    review_redirect_url,
    review_row_anchor,
)


def test_sanitize_filename_removes_paths_and_unsafe_characters() -> None:
    assert sanitize_filename("../bank statement июнь.pdf") == "bank_statement_.pdf"
    assert sanitize_filename("statement") == "statement.pdf"
    assert sanitize_upload_filename("../bank statement июнь.xlsx") == "bank_statement_.xlsx"


@pytest.mark.asyncio
async def test_upload_storage_preserves_pdf_bytes(tmp_path: Path) -> None:
    content = b"%PDF-1.4 local fixture bytes"
    upload = UploadFile(file=BytesIO(content), filename="../statement.pdf")
    workspace_id = uuid4()
    document_id = uuid4()

    stored = await UploadStorage(tmp_path).save_pdf(
        upload,
        workspace_id=workspace_id,
        document_id=document_id,
    )

    assert stored.file_size_bytes == len(content)
    assert stored.sha256_hash == sha256(content).hexdigest()
    assert stored.path.read_bytes() == content
    assert stored.storage_key == f"{workspace_id}/{document_id}/statement.pdf"


@pytest.mark.asyncio
async def test_upload_storage_save_pdf_keeps_legacy_pdf_suffix(tmp_path: Path) -> None:
    upload = UploadFile(file=BytesIO(b"%PDF-1.4"), filename="statement")
    workspace_id = uuid4()
    document_id = uuid4()

    stored = await UploadStorage(tmp_path).save_pdf(
        upload,
        workspace_id=workspace_id,
        document_id=document_id,
    )

    assert stored.storage_key == f"{workspace_id}/{document_id}/statement.pdf"


@pytest.mark.asyncio
async def test_upload_storage_preserves_xlsx_extension(tmp_path: Path) -> None:
    content = b"local xlsx fixture bytes"
    upload = UploadFile(file=BytesIO(content), filename="../statement.xlsx")
    workspace_id = uuid4()
    document_id = uuid4()

    stored = await UploadStorage(tmp_path).save_upload(
        upload,
        workspace_id=workspace_id,
        document_id=document_id,
    )

    assert stored.file_size_bytes == len(content)
    assert stored.sha256_hash == sha256(content).hexdigest()
    assert stored.path.read_bytes() == content
    assert stored.storage_key == f"{workspace_id}/{document_id}/statement.xlsx"


def test_validate_statement_upload_accepts_pdf_and_xlsx() -> None:
    validate_statement_upload(UploadFile(file=BytesIO(b"%PDF-1.4"), filename="statement.pdf"))
    validate_statement_upload(UploadFile(file=BytesIO(b"xlsx"), filename="statement.xlsx"))


def test_validate_statement_upload_rejects_unknown_extension() -> None:
    upload = UploadFile(file=BytesIO(b"not a pdf"), filename="statement.txt")

    with pytest.raises(UploadValidationError):
        validate_statement_upload(upload)


def test_pdfplumber_extractor_preserves_raw_pages() -> None:
    extracted = PdfPlumberExtractor().extract(Path("tests/fixtures/expobank_statement.pdf"))

    assert extracted.text_by_page
    assert len(extracted.tables_by_page) == len(extracted.text_by_page)
    assert all(page.page_number >= 1 for page in extracted.tables_by_page)


def test_openpyxl_extractor_preserves_sheet_tables(tmp_path: Path) -> None:
    workbook_path = tmp_path / "statement.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Card"
    sheet.append(["Дата", "Описание", "Сумма"])
    sheet.append(["2026-06-01", "Coffee", -10.5])
    sheet.append([None, None, None])
    workbook.save(workbook_path)

    extracted = OpenPyxlStatementExtractor().extract(workbook_path)

    assert extracted.metadata["source_format"] == "xlsx"
    assert extracted.metadata["sheet_names"] == ["Card"]
    assert extracted.text_by_page == ["Дата\tОписание\tСумма\n2026-06-01\tCoffee\t-10.5"]
    assert extracted.tables_by_page[0].page_number == 1
    assert extracted.tables_by_page[0].tables == [
        [
            ["Дата", "Описание", "Сумма"],
            ["2026-06-01", "Coffee", "-10.5"],
        ]
    ]


def test_statement_extractor_resolver_selects_extractor_by_extension(tmp_path: Path) -> None:
    workbook_path = tmp_path / "statement.xlsx"
    workbook = Workbook()
    workbook.active.append(["Date", "Amount"])
    workbook.save(workbook_path)

    extracted = StatementExtractorResolver().extract(workbook_path)

    assert extracted.metadata["source_format"] == "xlsx"


def test_raw_transaction_review_actions_map_to_statuses() -> None:
    assert raw_transaction_status_for_review_action("ignore") == RawTransactionStatus.IGNORED
    assert raw_transaction_status_for_review_action("mark_unique") == RawTransactionStatus.MATCHED
    assert (
        raw_transaction_status_for_review_action("needs_review")
        == RawTransactionStatus.NEEDS_REVIEW
    )


def test_possible_duplicate_fingerprint_requires_normalized_fields() -> None:
    account_id = uuid4()
    raw_transaction = raw_transaction_from_values(
        account_id=account_id,
        amount=Decimal("10.00"),
    )

    assert possible_duplicate_fingerprint(raw_transaction) == (
        account_id,
        parse_bank_date("29.05.2026"),
        Decimal("10.00"),
        "RUB",
    )
    raw_transaction.amount = None
    assert possible_duplicate_fingerprint(raw_transaction) is None


def test_mark_raw_transaction_duplicate_preserves_review_message() -> None:
    raw_transaction = raw_transaction_from_values(normalization_error="Existing issue.")

    mark_raw_transaction_duplicate(
        raw_transaction,
        RawTransactionStatus.DUPLICATE,
        "Exact duplicate.",
    )

    assert raw_transaction.status == RawTransactionStatus.DUPLICATE
    assert raw_transaction.normalization_error == "Existing issue.; Exact duplicate."


def test_document_has_linked_operations_detects_confirmed_rows() -> None:
    raw_transaction = raw_transaction_from_values()
    document = RawTransactionDocumentStub(raw_transactions=[raw_transaction])

    assert document_has_linked_operations(document) is False

    raw_transaction.linked_operation_id = uuid4()
    assert document_has_linked_operations(document) is True


def test_review_redirect_url_keeps_user_at_raw_transaction() -> None:
    document_id = uuid4()
    raw_transaction_id = uuid4()

    assert review_row_anchor(raw_transaction_id) == f"raw-{raw_transaction_id}"
    assert review_redirect_url(document_id) == f"/imports/documents/{document_id}/review"
    assert review_redirect_url(document_id, raw_transaction_id) == (
        f"/imports/documents/{document_id}/review#raw-{raw_transaction_id}"
    )


def raw_transaction_from_values(
    *,
    account_id: UUID | None = None,
    amount: Decimal | None = Decimal("10.00"),
    normalization_error: str | None = None,
) -> RawTransaction:
    return RawTransaction(
        workspace_id=uuid4(),
        uploaded_document_id=uuid4(),
        parse_attempt_id=uuid4(),
        row_index=0,
        status=RawTransactionStatus.NORMALIZED,
        raw_payload={},
        account_id=account_id,
        operation_date=parse_bank_date("29.05.2026"),
        amount=amount,
        currency="RUB",
        normalization_error=normalization_error,
    )


@dataclass(frozen=True)
class RawTransactionDocumentStub:
    raw_transactions: list[RawTransaction]
